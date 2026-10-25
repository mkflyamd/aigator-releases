"""LibreOffice discovery, macro runner and the Excel recalculation built on them."""
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

SKILLS = Path(__file__).resolve().parent.parent / "web" / "skills"
SHARED_SCRIPTS = SKILLS / "_scripts"
if str(SHARED_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SHARED_SCRIPTS))

from office import soffice  # noqa: E402
from skills.excel import tools as xl  # noqa: E402


@pytest.fixture
def recalc_mod(monkeypatch):
    mod = xl._load_recalc()
    monkeypatch.setattr(mod, "_excel_recalc", lambda *a, **k: False)
    return mod


def _fake_proc(returncode=0, stdout="", stderr=""):
    return subprocess.CompletedProcess([], returncode, stdout, stderr)


def _workbook_with_cached_values(path, values):
    """openpyxl writes formulas without results; add the cached values LibreOffice would write."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["A2"] = 3
    ws["A3"] = "=A1*A2"
    ws["A4"] = "=1/0"
    wb.save(path)
    with zipfile.ZipFile(path) as zf:
        members = {name: zf.read(name) for name in zf.namelist()}
    sheet = members["xl/worksheets/sheet1.xml"].decode("utf-8")
    sheet = sheet.replace('<c r="A3"><f>A1*A2</f><v></v></c>', f'<c r="A3"><f>A1*A2</f><v>{values["A3"]}</v></c>')
    sheet = sheet.replace('<c r="A4"><f>1/0</f><v></v></c>', '<c r="A4" t="e"><f>1/0</f><v>#DIV/0!</v></c>')
    members["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)


# ── discovery ────────────────────────────────────────────────────────────────


def test_find_soffice_prefers_path(monkeypatch):
    monkeypatch.setattr(soffice.shutil, "which", lambda name: "/usr/bin/soffice" if name == "soffice" else None)
    assert soffice.find_soffice() == "/usr/bin/soffice"


def test_find_soffice_windows_install_folder_not_on_path(monkeypatch, tmp_path):
    exe = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    exe.parent.mkdir(parents=True)
    exe.write_text("")
    monkeypatch.setattr(soffice.shutil, "which", lambda name: None)
    monkeypatch.setattr(soffice.sys, "platform", "win32")
    monkeypatch.setenv("ProgramFiles", str(tmp_path))
    monkeypatch.delenv("ProgramFiles(x86)", raising=False)
    monkeypatch.delenv("ProgramW6432", raising=False)
    monkeypatch.delenv("LOCALAPPDATA", raising=False)
    assert soffice.find_soffice() == str(exe)


def test_find_soffice_none_when_not_installed(monkeypatch):
    monkeypatch.setattr(soffice.shutil, "which", lambda name: None)
    monkeypatch.setattr(soffice, "_candidate_paths", lambda: [])
    assert soffice.find_soffice() is None
    with pytest.raises(soffice.SofficeNotFound):
        soffice.require_soffice()


def test_not_found_is_a_file_not_found_error():
    assert issubclass(soffice.SofficeNotFound, FileNotFoundError)


# ── environment ──────────────────────────────────────────────────────────────


def test_headless_backend_only_forced_off_windows(monkeypatch):
    monkeypatch.setattr(soffice, "_needs_shim", lambda: False)
    monkeypatch.delenv("SAL_USE_VCLPLUGIN", raising=False)
    monkeypatch.setattr(soffice.sys, "platform", "win32")
    assert "SAL_USE_VCLPLUGIN" not in soffice.get_soffice_env()
    monkeypatch.setattr(soffice.sys, "platform", "linux")
    assert soffice.get_soffice_env()["SAL_USE_VCLPLUGIN"] == "svp"


def test_socket_shim_is_linux_only(monkeypatch):
    for platform in ("win32", "darwin"):
        monkeypatch.setattr(soffice.sys, "platform", platform)
        assert soffice._needs_shim() is False


# ── macro runner ─────────────────────────────────────────────────────────────


def test_run_basic_macro_prepares_private_profile_once(monkeypatch, tmp_path):
    calls = []
    monkeypatch.setattr(soffice, "require_soffice", lambda: "fake-soffice")
    monkeypatch.setattr(soffice.tempfile, "gettempdir", lambda: str(tmp_path))
    monkeypatch.setattr(soffice, "_run", lambda cmd, timeout: calls.append((cmd, timeout)) or _fake_proc())
    doc = tmp_path / "book.xlsx"
    doc.write_text("")

    soffice.run_basic_macro("<macro Sub Do/>", "Do", str(doc), 20, "unit")
    soffice.run_basic_macro("<macro Sub Do/>", "Do", str(doc), 20, "unit")

    init_calls = [c for c in calls if "--terminate_after_init" in c[0]]
    macro_calls = [c for c in calls if any("Standard.Module1.Do" in a for a in c[0])]
    assert len(init_calls) == 1 and len(macro_calls) == 2
    profile = tmp_path / "gator_lo_unit_profile"
    assert (profile / "user" / "basic" / "Standard" / "Module1.xba").read_text(encoding="utf-8") == "<macro Sub Do/>"
    assert f"-env:UserInstallation={profile.as_uri()}" in macro_calls[0][0]
    assert macro_calls[0][0][0] == "fake-soffice"
    assert macro_calls[0][0][-1] == str(doc.absolute())
    assert macro_calls[0][1] == 20


def test_run_basic_macro_without_libreoffice(monkeypatch):
    monkeypatch.setattr(soffice.shutil, "which", lambda name: None)
    monkeypatch.setattr(soffice, "_candidate_paths", lambda: [])
    with pytest.raises(soffice.SofficeNotFound):
        soffice.run_basic_macro("x", "Do", "book.xlsx", 5, "unit")


def test_run_times_out_and_stops_the_process():
    with pytest.raises(subprocess.TimeoutExpired):
        soffice._run([sys.executable, "-c", "import time; time.sleep(60)"], 1)


# ── recalc ───────────────────────────────────────────────────────────────────


def test_recalc_reports_missing_libreoffice(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    Workbook().save(book)
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: None)
    out = recalc_mod.recalc(str(book))
    assert out["libreoffice_missing"] is True
    assert "not installed" in out["error"] and "Excel" in out["error"]


def test_recalc_reports_broken_libreoffice_install(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    Workbook().save(book)
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: "soffice")
    stderr = "The application cannot be started.\nThe component manager is not available.\n"
    monkeypatch.setattr(recalc_mod, "run_basic_macro", lambda *a, **k: _fake_proc(3221225477, "", stderr))
    out = recalc_mod.recalc(str(book))
    assert "installed but failed to run" in out["error"] and "3221225477" in out["error"]
    assert "component manager is not available" in out["error"]
    assert "libreoffice_missing" not in out


def test_recalc_reports_timeout(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    Workbook().save(book)
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: "soffice")

    def hang(*a, **k):
        raise subprocess.TimeoutExpired("soffice", 5)

    monkeypatch.setattr(recalc_mod, "run_basic_macro", hang)
    assert "did not finish within 7 seconds" in recalc_mod.recalc(str(book), 7)["error"]


def test_recalc_scans_calculated_workbook(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    _workbook_with_cached_values(book, {"A3": 6})
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(recalc_mod, "run_basic_macro", lambda *a, **k: _fake_proc())
    out = recalc_mod.recalc(str(book))
    assert out["status"] == "errors_found"
    assert out["total_formulas"] == 2
    assert out["error_summary"]["#DIV/0!"]["locations"] == ["Sheet!A4"]


def test_recalc_does_not_claim_success_when_nothing_was_calculated(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=1+1"
    wb.save(book)
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(recalc_mod, "run_basic_macro", lambda *a, **k: _fake_proc())
    assert "did not calculate any formulas" in recalc_mod.recalc(str(book))["error"]


def test_recalc_uses_excel_when_it_does_the_work(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    _workbook_with_cached_values(book, {"A3": 6})
    monkeypatch.setattr(recalc_mod, "_excel_recalc", lambda *a, **k: True)

    def no_libreoffice(*a, **k):
        raise AssertionError("LibreOffice must not run when Excel recalculated")

    monkeypatch.setattr(recalc_mod, "find_soffice", no_libreoffice)
    monkeypatch.setattr(recalc_mod, "run_basic_macro", no_libreoffice)
    out = recalc_mod.recalc(str(book))
    assert out["engine"] == "Excel"
    assert out["error_summary"]["#DIV/0!"]["locations"] == ["Sheet!A4"]


def test_recalc_falls_back_to_libreoffice_when_excel_cannot_be_used(recalc_mod, monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    _workbook_with_cached_values(book, {"A3": 6})
    monkeypatch.setattr(recalc_mod, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(recalc_mod, "run_basic_macro", lambda *a, **k: _fake_proc())
    assert recalc_mod.recalc(str(book))["engine"] == "LibreOffice"


def test_excel_recalc_is_skipped_off_windows(monkeypatch, tmp_path):
    mod = xl._load_recalc()
    monkeypatch.setattr(mod.sys, "platform", "linux")
    assert mod._excel_recalc(str(tmp_path / "book.xlsx"), 5) is False


def test_excel_recalc_kills_excel_when_it_hangs(monkeypatch, tmp_path):
    pythoncom = pytest.importorskip("pythoncom")
    win32com_client = pytest.importorskip("win32com.client")
    mod = xl._load_recalc()
    monkeypatch.setattr(mod.sys, "platform", "win32")
    book = tmp_path / "book.xlsx"
    Workbook().save(book)
    release = mod.threading.Event()

    class HungApp:
        Hwnd = 1
        Visible = False
        DisplayAlerts = False

        @property
        def Workbooks(self):
            release.wait(10)
            raise RuntimeError("released")

        def Quit(self):
            pass

    killed = []
    monkeypatch.setattr(win32com_client, "DispatchEx", lambda name: HungApp())
    monkeypatch.setattr(mod, "_excel_pid", lambda app: 4242)
    monkeypatch.setattr(mod.subprocess, "run", lambda cmd, **k: killed.append(cmd))
    try:
        assert mod._excel_recalc(str(book), 0.3) is False
    finally:
        release.set()
    assert killed and killed[0][:3] == ["taskkill", "/PID", "4242"]


def test_excel_recalc_ignores_a_save_that_did_not_change_the_file(monkeypatch, tmp_path):
    pytest.importorskip("pythoncom")
    win32com_client = pytest.importorskip("win32com.client")
    mod = xl._load_recalc()
    monkeypatch.setattr(mod.sys, "platform", "win32")
    book = tmp_path / "book.xlsx"
    Workbook().save(book)

    class Book:
        ReadOnly = False

        def Save(self):
            pass

        def Close(self, save):
            pass

    class SilentApp:
        Hwnd = 1
        Visible = False
        DisplayAlerts = False
        Workbooks = type("W", (), {"Open": staticmethod(lambda *a, **k: Book())})

        def CalculateFull(self):
            pass

        def Quit(self):
            pass

    monkeypatch.setattr(win32com_client, "DispatchEx", lambda name: SilentApp())
    monkeypatch.setattr(mod, "_excel_pid", lambda app: 1)
    assert mod._excel_recalc(str(book), 5) is False


def test_recalc_tool_passes_the_missing_libreoffice_message_through(monkeypatch, tmp_path):
    book = tmp_path / "book.xlsx"
    Workbook().save(book)
    monkeypatch.setattr(soffice.shutil, "which", lambda name: None)
    monkeypatch.setattr(soffice, "_candidate_paths", lambda: [])
    out = xl._tool_recalc_excel(str(book))
    assert out["libreoffice_missing"] is True
    assert "not installed" in out["error"]


def test_recalc_tool_summarises_a_successful_run(monkeypatch):
    class Stub:
        @staticmethod
        def recalc(path, timeout):
            return {"status": "success", "total_errors": 0, "total_formulas": 4, "error_summary": {}, "engine": "Excel"}

    monkeypatch.setattr(xl, "_load_recalc", lambda: Stub)
    assert xl._tool_recalc_excel("book.xlsx") == {
        "ok": True, "status": "success", "total_formulas": 4, "error_count": 0, "error_summary": {},
        "engine": "Excel",
    }


# ── the scripts run on their own, from any folder ────────────────────────────


@pytest.mark.parametrize(
    "script, args, expected",
    [
        ("excel/scripts/recalc.py", ["missing.xlsx"], "does not exist"),
        ("docx/scripts/accept_changes.py", ["missing.docx", "out.docx"], "Input file not found"),
        ("ppt/scripts/thumbnail.py", ["missing.pptx"], "Invalid PowerPoint file"),
    ],
)
def test_office_scripts_import_their_helpers_from_any_folder(script, args, expected, tmp_path):
    run = subprocess.run(
        [sys.executable, str(SKILLS / script), *args],
        cwd=tmp_path, capture_output=True, text=True, timeout=60,
    )
    assert "ModuleNotFoundError" not in run.stderr
    assert expected in run.stdout + run.stderr


def test_run_basic_macro_stops_when_libreoffice_cannot_start(monkeypatch, tmp_path):
    calls = []
    monkeypatch.setattr(soffice, "require_soffice", lambda: "fake-soffice")
    monkeypatch.setattr(soffice.tempfile, "gettempdir", lambda: str(tmp_path))
    monkeypatch.setattr(soffice, "_run", lambda cmd, timeout: calls.append(cmd) or _fake_proc(3, "", "boom"))
    result = soffice.run_basic_macro("<m/>", "Do", str(tmp_path / "b.xlsx"), 20, "unit")
    assert result.returncode == 3 and len(calls) == 1
